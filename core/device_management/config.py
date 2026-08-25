from __future__ import annotations

import re
from pathlib import Path
from typing import Mapping

from openpyxl import load_workbook

from .models import (
    ChannelFilters,
    DeviceConfig,
    DriverConfigSpec,
    ProfileConfig,
    ProfilePaths,
)


SCHEMA_VERSION = 1
DEVICE_SHEET_NAME = "Devices"
DEVICE_HEADERS = (
    "id",
    "driver",
    "enabled",
    "connect_on_start",
    "address",
    "scan set",
    "scan get",
    "config check",
)
DEVICE_ID_PATTERN = re.compile(r"^[A-Za-z0-9][A-Za-z0-9_-]*$")
RESERVED_DEVICE_IDS = frozenset({"artificial_channel", "default"})


class ProfileValidationError(ValueError):
    """Raised only after all safely collectible profile errors are gathered."""

    def __init__(self, errors: list[str] | tuple[str, ...]):
        self.errors = tuple(errors)
        detail = "\n".join(f"- {error}" for error in self.errors)
        super().__init__(f"Invalid ZMeter profile:\n{detail}")


def repository_root_from_module() -> Path:
    return Path(__file__).resolve().parents[2]


def _is_blank(value: object) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _parse_device_id(value: object, *, row_number: int, errors: list[str]) -> str:
    context = f"{DEVICE_SHEET_NAME}!A{row_number} (id)"
    if not isinstance(value, str) or not value.strip():
        errors.append(f"{context} must be a non-empty text value")
        return ""

    device_id = value.strip()
    if DEVICE_ID_PATTERN.fullmatch(device_id) is None:
        errors.append(
            f"{context} '{device_id}' may contain only letters, digits, "
            "underscores, and hyphens"
        )
    if device_id in RESERVED_DEVICE_IDS:
        errors.append(f"{context} '{device_id}' is reserved")
    return device_id


def _canonical_driver_ids(
    driver_specs: Mapping[str, DriverConfigSpec],
) -> dict[str, str]:
    canonical: dict[str, str] = {}
    for registry_id, spec in driver_specs.items():
        for candidate in (registry_id, spec.driver_id):
            normalized = candidate.strip().casefold()
            if normalized:
                canonical.setdefault(normalized, spec.driver_id)
    return canonical


def _parse_driver_id(
    value: object,
    *,
    row_number: int,
    canonical_driver_ids: Mapping[str, str],
    errors: list[str],
) -> str:
    context = f"{DEVICE_SHEET_NAME}!B{row_number} (driver)"
    if not isinstance(value, str) or not value.strip():
        errors.append(f"{context} must be a non-empty text value")
        return ""

    configured = value.strip()
    driver_id = canonical_driver_ids.get(configured.casefold())
    if driver_id is None:
        errors.append(f"{context} '{configured}' is not registered")
        return configured.casefold()
    return driver_id


def _parse_flag(
    value: object,
    *,
    row_number: int,
    column: str,
    field_name: str,
    errors: list[str],
) -> bool:
    if type(value) is bool:
        return value
    if isinstance(value, str):
        normalized = value.strip().casefold()
        if normalized == "true":
            return True
        if normalized == "false":
            return False

    errors.append(
        f"{DEVICE_SHEET_NAME}!{column}{row_number} ({field_name}) "
        "must be TRUE or FALSE"
    )
    return False


def _parse_address(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, str):
        return value
    # Address content is deliberately not validated. Converting a legacy
    # numeric cell keeps startup non-fatal while the workbook's Text format
    # prevents new identifiers from losing leading zeroes.
    return str(value)


def _parse_channel_filter(
    value: object,
    *,
    row_number: int,
    column: str,
    field_name: str,
    errors: list[str],
) -> tuple[str, ...] | None:
    if _is_blank(value):
        return None
    if not isinstance(value, str):
        errors.append(
            f"{DEVICE_SHEET_NAME}!{column}{row_number} ({field_name}) "
            "must be comma-separated text or blank"
        )
        return None

    channels = tuple(part.strip() for part in value.split(",") if part.strip())
    return channels or None


def _parse_devices(
    worksheet,
    *,
    driver_specs: Mapping[str, DriverConfigSpec],
) -> tuple[tuple[DeviceConfig, ...], list[str]]:
    errors: list[str] = []
    devices: list[DeviceConfig] = []
    seen_ids: dict[str, int] = {}
    canonical_driver_ids = _canonical_driver_ids(driver_specs)

    for row_number, row in enumerate(
        worksheet.iter_rows(
            min_row=2,
            min_col=1,
            max_col=len(DEVICE_HEADERS),
            values_only=True,
        ),
        start=2,
    ):
        # Reserved rows may contain a config-check formula in column H. They
        # are not device entries until at least one input cell A:G is filled.
        if all(_is_blank(value) for value in row[:7]):
            continue

        device_id = _parse_device_id(row[0], row_number=row_number, errors=errors)
        driver_id = _parse_driver_id(
            row[1],
            row_number=row_number,
            canonical_driver_ids=canonical_driver_ids,
            errors=errors,
        )
        enabled = _parse_flag(
            row[2],
            row_number=row_number,
            column="C",
            field_name="enabled",
            errors=errors,
        )
        connect_on_start = _parse_flag(
            row[3],
            row_number=row_number,
            column="D",
            field_name="connect_on_start",
            errors=errors,
        )

        if connect_on_start and not enabled:
            errors.append(
                f"{DEVICE_SHEET_NAME}!D{row_number} (connect_on_start) cannot "
                "be TRUE when enabled is FALSE"
        )

        if device_id:
            normalized_device_id = device_id.casefold()
            previous_row = seen_ids.get(normalized_device_id)
            if previous_row is None:
                seen_ids[normalized_device_id] = row_number
            else:
                errors.append(
                    f"{DEVICE_SHEET_NAME}!A{row_number} duplicates "
                    f"{DEVICE_SHEET_NAME}!A{previous_row} id '{device_id}'"
                )

        setters = _parse_channel_filter(
            row[5],
            row_number=row_number,
            column="F",
            field_name="scan set",
            errors=errors,
        )
        getters = _parse_channel_filter(
            row[6],
            row_number=row_number,
            column="G",
            field_name="scan get",
            errors=errors,
        )

        devices.append(
            DeviceConfig(
                id=device_id,
                driver=driver_id,
                enabled=enabled,
                connect_on_start=connect_on_start,
                connection={"address": _parse_address(row[4])},
                scan_channels=ChannelFilters(setters=setters, getters=getters),
            )
        )

    return tuple(devices), errors


def _profile_from_workbook(
    workbook,
    *,
    source: Path,
    repository_root: Path,
    driver_specs: Mapping[str, DriverConfigSpec],
) -> ProfileConfig:
    if DEVICE_SHEET_NAME not in workbook.sheetnames:
        raise ProfileValidationError(
            [f"workbook must contain a '{DEVICE_SHEET_NAME}' sheet"]
        )

    worksheet = workbook[DEVICE_SHEET_NAME]
    header_row = next(
        worksheet.iter_rows(
            min_row=1,
            max_row=1,
            min_col=1,
            max_col=len(DEVICE_HEADERS),
            values_only=True,
        ),
        (),
    )
    if tuple(header_row) != DEVICE_HEADERS:
        expected = ", ".join(DEVICE_HEADERS)
        found = ", ".join(
            "<blank>" if value is None else str(value) for value in header_row
        )
        raise ProfileValidationError(
            [
                f"{DEVICE_SHEET_NAME} headers A1:H1 must be exactly: "
                f"{expected}; found: {found or '<empty row>'}"
            ]
        )

    devices, errors = _parse_devices(worksheet, driver_specs=driver_specs)
    if errors:
        raise ProfileValidationError(errors)

    return ProfileConfig(
        schema_version=SCHEMA_VERSION,
        profile=source.stem,
        paths=ProfilePaths(save=repository_root / "data", backup=None),
        devices=devices,
        source_path=source,
        repository_root=repository_root,
    )


def load_profile(
    profile_path: str | Path,
    *,
    driver_specs: Mapping[str, DriverConfigSpec],
    repository_root: str | Path | None = None,
) -> ProfileConfig:
    """Load one Excel profile without importing or constructing a driver."""

    root = Path(repository_root or repository_root_from_module()).resolve()
    source = Path(profile_path)
    if not source.is_absolute():
        source = root / source
    source = source.resolve(strict=False)

    if source.suffix.casefold() != ".xlsx":
        raise ProfileValidationError(
            [f"profile must be an .xlsx workbook: {source}"]
        )

    workbook = None
    try:
        workbook = load_workbook(source, read_only=True, data_only=False)
        return _profile_from_workbook(
            workbook,
            source=source,
            repository_root=root,
            driver_specs=driver_specs,
        )
    except ProfileValidationError:
        raise
    except FileNotFoundError as exc:
        raise ProfileValidationError([f"profile file not found: {source}"]) from exc
    except Exception as exc:
        raise ProfileValidationError(
            [f"profile workbook could not be read: {type(exc).__name__}: {exc}"]
        ) from exc
    finally:
        if workbook is not None:
            workbook.close()
