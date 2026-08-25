from __future__ import annotations

from pathlib import Path
from typing import Iterable, Mapping

from openpyxl import Workbook


DEVICE_HEADERS = (
    "id",
    "driver",
    "enabled",
    "connect_on_start",
    "address",
    "scan set",
    "scan get",
    "config check",
)

MOCK_DEVICE_ROWS = (
    {
        "id": "mock_device_1",
        "driver": "mock_device",
        "enabled": True,
        "connect_on_start": False,
        "address": "MOCK::INSTR",
    },
    {
        "id": "mock_device_2",
        "driver": "mock_device",
        "enabled": True,
        "connect_on_start": False,
        "address": "MOCK::INSTR",
    },
)


def write_test_device_config(
    path: str | Path,
    rows: Iterable[Mapping[str, object]] = MOCK_DEVICE_ROWS,
) -> Path:
    """Write a minimal hardware-independent Devices workbook for a test."""

    destination = Path(path)
    destination.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Devices"
    worksheet.append(DEVICE_HEADERS)
    for row in rows:
        worksheet.append(
            (
                row.get("id"),
                row.get("driver"),
                row.get("enabled", False),
                row.get("connect_on_start", False),
                row.get("address", ""),
                row.get("scan set"),
                row.get("scan get"),
                row.get("config check"),
            )
        )
        worksheet.cell(worksheet.max_row, 5).number_format = "@"
    workbook.save(destination)
    workbook.close()
    return destination
