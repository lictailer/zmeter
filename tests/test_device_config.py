from __future__ import annotations

import subprocess
import sys
import tempfile
import unittest
from dataclasses import FrozenInstanceError
from pathlib import Path
from unittest import mock

from openpyxl import Workbook

from core.device_management.config import (
    DEVICE_HEADERS,
    ProfileValidationError,
    load_profile,
)
from core.device_management.models import ConnectionFieldSpec, DriverConfigSpec


class DeviceProfileConfigTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name).resolve()
        self.driver_specs = {
            "mock_device": DriverConfigSpec(
                driver_id="mock_device",
                connection_fields={"address": ConnectionFieldSpec((str,))},
            ),
            "unavailable_driver": DriverConfigSpec(
                driver_id="unavailable_driver",
                connection_fields={},
                available=False,
                unavailable_reason="optional SDK is missing",
            ),
        }

    @staticmethod
    def valid_row(**overrides):
        values = {
            "id": "mock_device_1",
            "driver": "mock_device",
            "enabled": True,
            "connect_on_start": False,
            "address": "00000000",
            "scan_set": "set_channel_A, future_channel",
            "scan_get": "get_channel_A",
            "config_check": '=IF(A2="","MISSING ID","ENABLED")',
        }
        values.update(overrides)
        return [
            values["id"],
            values["driver"],
            values["enabled"],
            values["connect_on_start"],
            values["address"],
            values["scan_set"],
            values["scan_get"],
            values["config_check"],
        ]

    def write_workbook(
        self,
        rows=None,
        *,
        name="test.xlsx",
        headers=DEVICE_HEADERS,
        sheet_name="Devices",
        formula_only_row=False,
    ):
        path = self.root / name
        path.parent.mkdir(parents=True, exist_ok=True)
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = sheet_name
        worksheet.append(list(headers))
        for row in rows if rows is not None else [self.valid_row()]:
            worksheet.append(row)
        if formula_only_row:
            worksheet.cell(
                row=worksheet.max_row + 1,
                column=8,
                value='=IF(A3="","","ENABLED")',
            )
        workbook.save(path)
        workbook.close()
        return path

    def load(self, path):
        return load_profile(
            path,
            driver_specs=self.driver_specs,
            repository_root=self.root,
        )

    def test_valid_workbook_builds_immutable_profile_with_fixed_paths(self):
        path = self.write_workbook(name="lab_profile.xlsx")

        profile = self.load(path)

        self.assertEqual(profile.schema_version, 1)
        self.assertEqual(profile.profile, "lab_profile")
        self.assertEqual(profile.paths.save, self.root / "data")
        self.assertIsNone(profile.paths.backup)
        self.assertFalse((self.root / "data").exists())
        self.assertEqual(profile.source_path, path)
        self.assertEqual([device.id for device in profile.devices], ["mock_device_1"])
        device = profile.devices[0]
        self.assertEqual(device.connection["address"], "00000000")
        self.assertEqual(
            device.scan_channels.setters,
            ("set_channel_A", "future_channel"),
        )
        self.assertEqual(device.scan_channels.getters, ("get_channel_A",))
        with self.assertRaises(TypeError):
            device.connection["address"] = "changed"
        with self.assertRaises(FrozenInstanceError):
            profile.profile = "changed"

    def test_relative_workbook_path_does_not_depend_on_current_directory(self):
        path = self.write_workbook(name="configs/relative.xlsx")
        unrelated = self.root / "unrelated"
        unrelated.mkdir()

        with mock.patch("os.getcwd", return_value=str(unrelated)):
            profile = self.load(Path("configs") / "relative.xlsx")

        self.assertEqual(profile.source_path, path)

    def test_loading_in_fresh_process_does_not_import_device_or_vendor_modules(self):
        path = self.write_workbook()
        repository_root = Path(__file__).resolve().parents[1]
        script = """
import sys
from pathlib import Path
from core.device_management.config import load_profile
from core.device_management.models import ConnectionFieldSpec, DriverConfigSpec

specs = {
    "mock_device": DriverConfigSpec(
        driver_id="mock_device",
        connection_fields={"address": ConnectionFieldSpec((str,))},
    )
}
load_profile(Path(sys.argv[1]), driver_specs=specs, repository_root=Path(sys.argv[2]))
watched = ("mockDevice", "devices", "pyvisa", "clr")
loaded = sorted(
    name for name in sys.modules
    if any(name == prefix or name.startswith(prefix + ".") for prefix in watched)
)
if loaded:
    raise SystemExit("unexpected device/vendor imports: " + ", ".join(loaded))
print("profile load remained device/vendor-import free")
"""

        result = subprocess.run(
            [sys.executable, "-B", "-c", script, str(path), str(self.root)],
            cwd=repository_root,
            capture_output=True,
            text=True,
            timeout=10,
            check=False,
        )

        self.assertEqual(result.returncode, 0, result.stderr or result.stdout)
        self.assertIn("device/vendor-import free", result.stdout)

    def test_only_xlsx_files_are_accepted(self):
        json_path = self.root / "legacy.json"
        json_path.write_text("{}", encoding="utf-8")

        with self.assertRaisesRegex(ProfileValidationError, r"\.xlsx workbook"):
            self.load(json_path)
        with self.assertRaisesRegex(ProfileValidationError, "file not found"):
            self.load(self.root / "missing.xlsx")

        corrupt = self.root / "corrupt.xlsx"
        corrupt.write_bytes(b"not an Excel workbook")
        with self.assertRaisesRegex(ProfileValidationError, "could not be read"):
            self.load(corrupt)

    def test_devices_sheet_and_exact_headers_are_required(self):
        wrong_sheet = self.write_workbook(name="wrong-sheet.xlsx", sheet_name="Guide")
        with self.assertRaisesRegex(ProfileValidationError, "'Devices' sheet"):
            self.load(wrong_sheet)

        wrong_headers = list(DEVICE_HEADERS)
        wrong_headers[4] = "Address"
        path = self.write_workbook(name="wrong-headers.xlsx", headers=wrong_headers)
        with self.assertRaisesRegex(ProfileValidationError, "headers A1:H1 must be exactly"):
            self.load(path)

    def test_driver_names_are_case_insensitive_and_unknown_drivers_are_fatal(self):
        casefolded = self.load(
            self.write_workbook(
                [self.valid_row(driver="  MoCk_DeViCe  ")],
                name="casefolded.xlsx",
            )
        )
        self.assertEqual(casefolded.devices[0].driver, "mock_device")

        unknown = self.write_workbook(
            [self.valid_row(driver="MISSING_DRIVER", enabled=False)],
            name="unknown.xlsx",
        )
        with self.assertRaisesRegex(ProfileValidationError, "is not registered"):
            self.load(unknown)

    def test_known_unavailable_driver_does_not_block_profile_loading(self):
        profile = self.load(
            self.write_workbook(
                [self.valid_row(driver="UNAVAILABLE_DRIVER", enabled=True)],
                name="unavailable.xlsx",
            )
        )

        self.assertTrue(profile.devices[0].enabled)
        self.assertEqual(profile.devices[0].driver, "unavailable_driver")

    def test_native_and_text_booleans_are_accepted_case_insensitively(self):
        rows = [
            self.valid_row(enabled=True, connect_on_start="TrUe"),
            self.valid_row(
                id="mock_device_2",
                enabled=" false ",
                connect_on_start=False,
            ),
        ]

        profile = self.load(self.write_workbook(rows, name="booleans.xlsx"))

        self.assertTrue(profile.devices[0].enabled)
        self.assertTrue(profile.devices[0].connect_on_start)
        self.assertFalse(profile.devices[1].enabled)
        self.assertFalse(profile.devices[1].connect_on_start)

    def test_invalid_flags_and_connecting_a_disabled_device_are_collected(self):
        path = self.write_workbook(
            [
                self.valid_row(enabled=1, connect_on_start="sometimes"),
                self.valid_row(
                    id="mock_device_2",
                    enabled="FALSE",
                    connect_on_start="TRUE",
                ),
            ],
            name="bad-flags.xlsx",
        )

        with self.assertRaises(ProfileValidationError) as raised:
            self.load(path)

        message = str(raised.exception)
        self.assertIn("Devices!C2 (enabled) must be TRUE or FALSE", message)
        self.assertIn("Devices!D2 (connect_on_start) must be TRUE or FALSE", message)
        self.assertIn("cannot be TRUE when enabled is FALSE", message)

    def test_blank_duplicate_invalid_and_reserved_ids_are_rejected(self):
        path = self.write_workbook(
            [
                self.valid_row(id=""),
                self.valid_row(id="bad label"),
                self.valid_row(id="same_id"),
                self.valid_row(id="same_id"),
                self.valid_row(id="default"),
            ],
            name="bad-ids.xlsx",
        )

        with self.assertRaises(ProfileValidationError) as raised:
            self.load(path)

        message = str(raised.exception)
        self.assertIn("must be a non-empty text value", message)
        self.assertIn("may contain only", message)
        self.assertIn("duplicates Devices!A4", message)
        self.assertIn("is reserved", message)

    def test_addresses_are_never_structurally_validated(self):
        profile = self.load(
            self.write_workbook(
                [
                    self.valid_row(address=None),
                    self.valid_row(id="mock_device_2", address=35995686),
                    self.valid_row(
                        id="mock_device_3",
                        address="not:a:valid:endpoint",
                    ),
                ],
                name="addresses.xlsx",
            )
        )

        self.assertEqual(
            [device.connection["address"] for device in profile.devices],
            ["", "35995686", "not:a:valid:endpoint"],
        )

    def test_device_ids_are_unique_without_regard_to_case(self):
        path = self.write_workbook(
            [
                self.valid_row(id="device_one"),
                self.valid_row(id="DEVICE_ONE"),
            ],
            name="case-duplicate.xlsx",
        )

        with self.assertRaisesRegex(
            ProfileValidationError,
            r"Devices!A3 duplicates Devices!A2",
        ):
            self.load(path)

    def test_scan_channels_use_commas_and_non_text_values_are_rejected(self):
        profile = self.load(
            self.write_workbook(
                [
                    self.valid_row(
                        scan_set=" AO0, AO1, ,",
                        scan_get="AI9,AI12, counter0",
                    ),
                    self.valid_row(
                        id="mock_device_2",
                        scan_set=None,
                        scan_get="",
                    ),
                ],
                name="channels.xlsx",
            )
        )

        self.assertEqual(profile.devices[0].scan_channels.setters, ("AO0", "AO1"))
        self.assertEqual(
            profile.devices[0].scan_channels.getters,
            ("AI9", "AI12", "counter0"),
        )
        self.assertIsNone(profile.devices[1].scan_channels.setters)
        self.assertIsNone(profile.devices[1].scan_channels.getters)

        invalid = self.write_workbook(
            [self.valid_row(scan_set=4, scan_get=True)],
            name="invalid-channels.xlsx",
        )
        with self.assertRaises(ProfileValidationError) as raised:
            self.load(invalid)
        self.assertIn("Devices!F2 (scan set)", str(raised.exception))
        self.assertIn("Devices!G2 (scan get)", str(raised.exception))

    def test_config_check_values_are_ignored_and_formula_only_rows_are_skipped(self):
        profile = self.load(
            self.write_workbook(
                [self.valid_row(config_check="#VALUE!")],
                name="config-check.xlsx",
                formula_only_row=True,
            )
        )

        self.assertEqual([device.id for device in profile.devices], ["mock_device_1"])


if __name__ == "__main__":
    unittest.main()
