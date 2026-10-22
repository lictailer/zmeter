"""Read-only checks of the shipped configuration; never construct instruments."""
from pathlib import Path
import unittest

from openpyxl import load_workbook

from core.device_management import build_default_registry
from core.device_management.config import DEVICE_HEADERS, load_profile

WORKBOOK_PATH = Path(__file__).resolve().parents[1] / "device_config.xlsx"
ADDED_DRIVERS = ("attodry", "mc2000b", "spectraphysicsx3_laser", "auto_focus", "arduino_steppers", "auto_position")


class DeviceConfigWorkbookTests(unittest.TestCase):
    def setUp(self):
        self.workbook = load_workbook(WORKBOOK_PATH, data_only=False)
        self.addCleanup(self.workbook.close)
        self.devices = self.workbook["Devices"]

    def test_headers_table_extent_and_reserved_rows(self):
        self.assertEqual(self.workbook.sheetnames, ["Devices", "Guide"])
        self.assertEqual(tuple(c.value for c in self.devices[1]), DEVICE_HEADERS)
        self.assertEqual(self.devices.max_row, 201)
        self.assertEqual(list(self.devices.tables), ["FlatDeviceConfigTable"])
        self.assertEqual(self.devices.tables["FlatDeviceConfigTable"].ref, "A1:H30")
        self.assertEqual(self.devices.tables["FlatDeviceConfigTable"].tableStyleInfo.name, "TableStyleLight1")

    def test_existing_rows_keep_upstream_values_not_lab_autoconnect_settings(self):
        expected_ids = ["sr860_0", "sr860_1", "sr860_2", "sr830_0", "sr830_1", "sr830_2", "mfli_0",
                        "keithley24xx_0", "keithley24xx_1", "nidaq", "ni6423", "k10cr1_0", "k10cr1_1",
                        "pem100", "sp150", "hp34401a", "bbd30x", "four9", "montana2", "opticool", "tlpm",
                        "mock_device_0", "mock_device_1"]
        self.assertEqual([self.devices.cell(r, 1).value for r in range(2, 25)], expected_ids)
        for row in range(2, 25):
            self.assertIs(self.devices.cell(row, 3).value, row == 8)
            self.assertIs(self.devices.cell(row, 4).value, False)
        self.assertEqual(self.devices["E8"].value, "DEV30037, 192.168.141.54")
        self.assertEqual(self.devices["E9"].value, "GPIB1::18::INSTR")
        self.assertEqual(self.devices["F12"].value, "AO0, AO1")
        self.assertEqual(self.devices["G12"].value, "AI9,AI12, counter0")

    def test_new_devices_have_manual_disabled_rows_and_explicit_channels(self):
        profile = load_profile(WORKBOOK_PATH, driver_specs=build_default_registry().config_specs)
        configurations = {device.id: device for device in profile.devices}
        for row, driver in enumerate(ADDED_DRIVERS, 25):
            self.assertEqual([self.devices.cell(row, col).value for col in range(1, 5)],
                             [driver, driver, False, False])
            self.assertEqual(self.devices.cell(row, 5).number_format, "@")
            config = configurations[driver]
            self.assertFalse(config.enabled)
            self.assertFalse(config.connect_on_start)
            self.assertIsNotNone(config.scan_channels.setters)
            self.assertIsNotNone(config.scan_channels.getters)
        atto = configurations["attodry"]
        self.assertEqual(atto.connection["address"], "REPLACE_WITH_ATTODRY_ADDRESS")
        self.assertEqual(atto.scan_channels.setters, ("T", "T_stable", "B", "B_stable"))
        self.assertEqual(atto.scan_channels.getters, ("T", "B"))
        self.assertEqual(configurations["spectraphysicsx3_laser"].scan_channels.setters, ("wavelenght_stable", "shutter"))
        for driver in ("auto_focus", "auto_position"):
            self.assertIn(configurations[driver].connection["address"], (None, ""))
            self.assertIn("run_key", configurations[driver].scan_channels.getters)

    def test_guide_covers_every_registry_driver_and_alias(self):
        registry = build_default_registry()
        expected = set(registry.driver_ids) | {alias for spec in registry.config_specs.values() for alias in spec.aliases}
        accepted = [self.workbook["Guide"].cell(r, 4).value for r in range(4, 28)]
        self.assertEqual(set(accepted), expected)
        self.assertEqual(len(accepted), len(set(accepted)))
        for row in range(2, 202):
            formula = self.devices.cell(row, 8).value
            self.assertIn(f"COUNTA(A{row}:G{row})=0", formula)
            self.assertIn("$D$4:$D$27", formula)
            self.assertIn("UNKNOWN DRIVER", formula)
            self.assertIn("INVALID CONNECT FLAG", formula)

    def test_cached_checks_match_native_boolean_flags(self):
        cached = load_workbook(WORKBOOK_PATH, data_only=True)
        self.addCleanup(cached.close)
        for row in range(2, 31):
            self.assertIsInstance(self.devices.cell(row, 3).value, bool)
            self.assertIsInstance(self.devices.cell(row, 4).value, bool)
            self.assertEqual(cached["Devices"].cell(row, 8).value, "ENABLED" if row == 8 else "DISABLED")
        self.assertIn(cached["Devices"]["H201"].value, (None, ""))


if __name__ == "__main__":
    unittest.main()
